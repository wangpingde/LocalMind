"""文档解析."""

from __future__ import annotations

import csv
import json
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document as DocxDocument


class DocParser:
    SUPPORTED = {
        ".txt",
        ".md",
        ".pdf",
        ".docx",
        ".csv",
        ".xlsx",
        ".pptx",
        ".html",
        ".htm",
        ".json",
        ".py",
        ".java",
        ".sql",
        ".js",
        ".ts",
        ".yaml",
        ".yml",
    }

    def parse(self, file_path: Path) -> tuple[str, list[dict]]:
        """返回 (全文, 分段列表). 每段含 content, page_no, heading_path."""
        suffix = file_path.suffix.lower()
        parsers = {
            ".txt": self._parse_text,
            ".md": self._parse_text,
            ".pdf": self._parse_pdf,
            ".docx": self._parse_docx,
            ".csv": self._parse_csv,
            ".xlsx": self._parse_xlsx,
            ".pptx": self._parse_pptx,
            ".html": self._parse_html,
            ".htm": self._parse_html,
            ".json": self._parse_json,
            ".py": self._parse_code,
            ".java": self._parse_code,
            ".sql": self._parse_code,
            ".js": self._parse_code,
            ".ts": self._parse_code,
            ".yaml": self._parse_code,
            ".yml": self._parse_code,
        }
        handler = parsers.get(suffix)
        if not handler:
            raise ValueError(f"不支持的文件格式: {suffix}")
        return handler(file_path)

    def _parse_text(self, file_path: Path) -> tuple[str, list[dict]]:
        text = file_path.read_text(encoding="utf-8", errors="ignore")
        sections: list[dict] = []
        current_heading = ""
        buffer: list[str] = []

        for line in text.splitlines():
            if file_path.suffix.lower() == ".md" and line.startswith("#"):
                if buffer:
                    sections.append(
                        {
                            "content": "\n".join(buffer).strip(),
                            "page_no": None,
                            "heading_path": current_heading,
                        }
                    )
                    buffer = []
                current_heading = line.lstrip("#").strip()
            buffer.append(line)

        if buffer:
            sections.append(
                {
                    "content": "\n".join(buffer).strip(),
                    "page_no": None,
                    "heading_path": current_heading,
                }
            )

        if not sections and text.strip():
            sections.append({"content": text.strip(), "page_no": None, "heading_path": ""})

        return text, sections

    def _parse_pdf(self, file_path: Path) -> tuple[str, list[dict]]:
        doc = fitz.open(str(file_path))
        sections: list[dict] = []
        full_parts: list[str] = []
        for page_no, page in enumerate(doc, start=1):
            text = page.get_text()
            if text.strip():
                full_parts.append(text)
                sections.append(
                    {"content": text.strip(), "page_no": page_no, "heading_path": f"第{page_no}页"}
                )
        doc.close()
        return "\n\n".join(full_parts), sections

    def _parse_docx(self, file_path: Path) -> tuple[str, list[dict]]:
        doc = DocxDocument(str(file_path))
        sections: list[dict] = []
        current_heading = ""
        buffer: list[str] = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name or "").lower()
            if "heading" in style:
                if buffer:
                    sections.append(
                        {
                            "content": "\n".join(buffer),
                            "page_no": None,
                            "heading_path": current_heading,
                        }
                    )
                    buffer = []
                current_heading = text
            buffer.append(text)

        if buffer:
            sections.append(
                {"content": "\n".join(buffer), "page_no": None, "heading_path": current_heading}
            )

        full = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        return full, sections

    def _parse_csv(self, file_path: Path) -> tuple[str, list[dict]]:
        text = file_path.read_text(encoding="utf-8", errors="ignore")
        reader = csv.reader(text.splitlines())
        rows = list(reader)
        if not rows:
            return "", []

        sections: list[dict] = []
        header = rows[0]
        chunk_rows: list[str] = []
        chunk_size = 40

        for i in range(1, len(rows)):
            row = rows[i]
            chunk_rows.append(", ".join(f"{h}: {v}" for h, v in zip(header, row, strict=False)))
            if len(chunk_rows) >= chunk_size:
                sections.append(
                    {
                        "content": "\n".join(chunk_rows),
                        "page_no": None,
                        "heading_path": f"行 {i - chunk_size + 1}-{i}",
                    }
                )
                chunk_rows = []

        if chunk_rows:
            sections.append(
                {
                    "content": "\n".join(chunk_rows),
                    "page_no": None,
                    "heading_path": f"行 {max(1, len(rows) - len(chunk_rows))}-{len(rows) - 1}",
                }
            )

        full = "\n".join(",".join(r) for r in rows)
        return full, sections

    def _parse_xlsx(self, file_path: Path) -> tuple[str, list[dict]]:
        from openpyxl import load_workbook

        wb = load_workbook(str(file_path), read_only=True, data_only=True)
        sections: list[dict] = []
        full_parts: list[str] = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows: list[str] = []
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                if any(cells):
                    rows.append("\t".join(cells))
            if rows:
                content = "\n".join(rows)
                full_parts.append(f"[{sheet_name}]\n{content}")
                for i in range(0, len(rows), 50):
                    chunk = rows[i : i + 50]
                    sections.append(
                        {
                            "content": "\n".join(chunk),
                            "page_no": None,
                            "heading_path": f"{sheet_name} 行{i + 1}-{i + len(chunk)}",
                        }
                    )
        wb.close()
        return "\n\n".join(full_parts), sections

    def _parse_pptx(self, file_path: Path) -> tuple[str, list[dict]]:
        from pptx import Presentation

        prs = Presentation(str(file_path))
        sections: list[dict] = []
        full_parts: list[str] = []

        for idx, slide in enumerate(prs.slides, start=1):
            texts: list[str] = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text.strip())
            if texts:
                content = "\n".join(texts)
                full_parts.append(content)
                sections.append(
                    {
                        "content": content,
                        "page_no": idx,
                        "heading_path": f"幻灯片 {idx}",
                    }
                )

        return "\n\n".join(full_parts), sections

    def _parse_html(self, file_path: Path) -> tuple[str, list[dict]]:
        from bs4 import BeautifulSoup

        raw = file_path.read_text(encoding="utf-8", errors="ignore")
        soup = BeautifulSoup(raw, "html.parser")
        for tag in soup(["script", "style", "noscript"]):
            tag.decompose()
        text = soup.get_text("\n", strip=True)
        sections: list[dict] = []
        if text:
            for i in range(0, len(text), 3000):
                chunk = text[i : i + 3000]
                sections.append(
                    {
                        "content": chunk,
                        "page_no": None,
                        "heading_path": f"片段 {i // 3000 + 1}",
                    }
                )
        return text, sections

    def _parse_json(self, file_path: Path) -> tuple[str, list[dict]]:
        raw = file_path.read_text(encoding="utf-8", errors="ignore")
        data = json.loads(raw)
        pretty = json.dumps(data, ensure_ascii=False, indent=2)
        sections = [
            {
                "content": pretty[i : i + 3000],
                "page_no": None,
                "heading_path": f"JSON 片段 {i // 3000 + 1}",
            }
            for i in range(0, len(pretty), 3000)
        ] or [{"content": pretty, "page_no": None, "heading_path": "JSON"}]
        return pretty, sections

    def _parse_code(self, file_path: Path) -> tuple[str, list[dict]]:
        text = file_path.read_text(encoding="utf-8", errors="ignore")
        lines = text.splitlines()
        sections: list[dict] = []
        chunk_size = 80
        for i in range(0, len(lines), chunk_size):
            chunk = lines[i : i + chunk_size]
            sections.append(
                {
                    "content": "\n".join(chunk),
                    "page_no": None,
                    "heading_path": f"{file_path.name} L{i + 1}-{i + len(chunk)}",
                }
            )
        if not sections and text.strip():
            sections.append({"content": text, "page_no": None, "heading_path": file_path.name})
        return text, sections
