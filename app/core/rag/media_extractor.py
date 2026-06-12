"""从知识库文档中提取图片/视频帧等媒体资源."""

from __future__ import annotations

import hashlib
import re
import shutil
import subprocess
from dataclasses import dataclass, field
from pathlib import Path
from typing import Literal
from urllib.parse import urlparse

import fitz
from loguru import logger

from app.core.llm.vision_messages import MIME_BY_SUFFIX, guess_mime

MediaKind = Literal["image", "video_frame", "video_audio"]

IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}
VIDEO_SUFFIXES = {".mp4", ".mov", ".webm", ".mkv"}
MD_IMAGE_RE = re.compile(r"!\[[^\]]*\]\(([^)]+)\)")
HTML_IMG_SRC_RE = re.compile(r"""<img[^>]+src=['"]([^'"]+)['"]""", re.IGNORECASE)


@dataclass
class MediaAsset:
    kind: MediaKind
    source_path: str
    page_no: int | None
    heading_path: str
    mime: str
    data: bytes
    context_hint: str = ""
    content_hash: str = field(default="")

    def __post_init__(self) -> None:
        if not self.content_hash:
            self.content_hash = hashlib.sha256(self.data).hexdigest()

    @property
    def dedup_key(self) -> str:
        if self.data:
            return self.content_hash
        return f"{self.source_path}:{self.heading_path}:{self.context_hint}"


class MediaExtractor:
    def __init__(
        self,
        knowledge_dir: Path,
        *,
        video_max_frames: int = 12,
        video_frame_interval_sec: float = 30.0,
    ) -> None:
        self.knowledge_dir = knowledge_dir.resolve()
        self.video_max_frames = video_max_frames
        self.video_frame_interval_sec = video_frame_interval_sec

    def extract(
        self,
        file_path: Path,
        sections: list[dict] | None = None,
        *,
        rel_path: str | None = None,
    ) -> list[MediaAsset]:
        file_path = file_path.resolve()
        rel = rel_path or self._rel_path(file_path)
        sections = sections or []
        suffix = file_path.suffix.lower()

        if suffix in IMAGE_SUFFIXES:
            return self._extract_standalone_image(file_path, rel)
        if suffix in VIDEO_SUFFIXES:
            return self._extract_video(file_path, rel)

        extractors = {
            ".pdf": self._extract_pdf,
            ".docx": self._extract_docx,
            ".pptx": self._extract_pptx,
            ".html": self._extract_html,
            ".htm": self._extract_html,
            ".md": self._extract_markdown_images,
            ".txt": self._extract_markdown_images,
        }
        handler = extractors.get(suffix)
        if not handler:
            return []
        return handler(file_path, rel, sections)

    def _rel_path(self, file_path: Path) -> str:
        try:
            return file_path.relative_to(self.knowledge_dir).as_posix()
        except ValueError:
            return file_path.name

    def _context_for_page(self, sections: list[dict], page_no: int | None) -> str:
        if page_no is None:
            return self._nearest_context(sections, "")
        for sec in sections:
            if sec.get("page_no") == page_no:
                text = (sec.get("content") or "").strip()
                return text[:500] if text else ""
        return self._nearest_context(sections, "")

    def _nearest_context(self, sections: list[dict], heading: str) -> str:
        for sec in sections:
            if heading and sec.get("heading_path") == heading:
                return (sec.get("content") or "")[:500]
        if sections:
            return (sections[0].get("content") or "")[:500]
        return ""

    def _extract_standalone_image(self, file_path: Path, rel_path: str) -> list[MediaAsset]:
        data = file_path.read_bytes()
        if not data:
            return []
        return [
            MediaAsset(
                kind="image",
                source_path=rel_path,
                page_no=None,
                heading_path="[图片] 独立文件",
                mime=guess_mime(file_path),
                data=data,
                context_hint=file_path.name,
            )
        ]

    def _extract_video(self, file_path: Path, rel_path: str) -> list[MediaAsset]:
        ffmpeg = shutil.which("ffmpeg")
        if not ffmpeg:
            logger.info("未检测到 ffmpeg，视频 {} 仅记录元数据", rel_path)
            return [
                MediaAsset(
                    kind="video_frame",
                    source_path=rel_path,
                    page_no=None,
                    heading_path="[视频] 元数据",
                    mime="text/plain",
                    data=b"",
                    context_hint=(
                        f"视频文件 {file_path.name}（未安装 ffmpeg，无法抽帧分析）"
                    ),
                )
            ]

        frames = self._ffmpeg_extract_frames(ffmpeg, file_path)
        assets: list[MediaAsset] = []
        for idx, (timestamp_sec, frame_bytes) in enumerate(frames):
            assets.append(
                MediaAsset(
                    kind="video_frame",
                    source_path=rel_path,
                    page_no=None,
                    heading_path=f"[视频] {file_path.name} @ {timestamp_sec:.0f}s",
                    mime="image/jpeg",
                    data=frame_bytes,
                    context_hint=f"视频 {file_path.name} 第 {timestamp_sec:.0f} 秒画面",
                )
            )
        return assets

    def _ffmpeg_extract_frames(
        self, ffmpeg: str, file_path: Path
    ) -> list[tuple[float, bytes]]:
        duration = self._probe_duration(file_path)
        interval = max(1.0, self.video_frame_interval_sec)
        timestamps: list[float] = []
        t = 0.0
        while len(timestamps) < self.video_max_frames:
            if duration and t > duration:
                break
            timestamps.append(t)
            t += interval
        if not timestamps:
            timestamps = [0.0]

        frames: list[tuple[float, bytes]] = []
        for ts in timestamps:
            cmd = [
                ffmpeg,
                "-hide_banner",
                "-loglevel",
                "error",
                "-ss",
                str(ts),
                "-i",
                str(file_path),
                "-vframes",
                "1",
                "-f",
                "image2pipe",
                "-vcodec",
                "mjpeg",
                "pipe:1",
            ]
            try:
                proc = subprocess.run(cmd, capture_output=True, timeout=120, check=False)
                if proc.returncode == 0 and proc.stdout:
                    frames.append((ts, proc.stdout))
            except Exception as e:
                logger.warning("视频抽帧失败 {} @{}s: {}", file_path.name, ts, e)
        return frames

    def _probe_duration(self, file_path: Path) -> float | None:
        ffprobe = shutil.which("ffprobe")
        if not ffprobe:
            return None
        cmd = [
            ffprobe,
            "-v",
            "error",
            "-show_entries",
            "format=duration",
            "-of",
            "default=noprint_wrappers=1:nokey=1",
            str(file_path),
        ]
        try:
            proc = subprocess.run(cmd, capture_output=True, text=True, timeout=30, check=False)
            if proc.returncode == 0 and proc.stdout.strip():
                return float(proc.stdout.strip())
        except Exception:
            pass
        return None

    def _extract_pdf(
        self, file_path: Path, rel_path: str, sections: list[dict]
    ) -> list[MediaAsset]:
        assets: list[MediaAsset] = []
        doc = fitz.open(str(file_path))
        seen: set[str] = set()
        try:
            for page_no, page in enumerate(doc, start=1):
                context = self._context_for_page(sections, page_no)
                for img_idx, img_info in enumerate(page.get_images(full=True), start=1):
                    xref = img_info[0]
                    key = f"pdf:{xref}"
                    if key in seen:
                        continue
                    seen.add(key)
                    try:
                        extracted = doc.extract_image(xref)
                        data = extracted.get("image", b"")
                        if not data:
                            continue
                        ext = extracted.get("ext", "png")
                        mime = MIME_BY_SUFFIX.get(f".{ext}", "image/png")
                        assets.append(
                            MediaAsset(
                                kind="image",
                                source_path=rel_path,
                                page_no=page_no,
                                heading_path=f"[图片] 第{page_no}页/图{img_idx}",
                                mime=mime,
                                data=data,
                                context_hint=context,
                            )
                        )
                    except Exception as e:
                        logger.debug("PDF 图片提取失败 {} p{}: {}", rel_path, page_no, e)
        finally:
            doc.close()
        return assets

    def _extract_docx(
        self, file_path: Path, rel_path: str, sections: list[dict]
    ) -> list[MediaAsset]:
        from docx import Document as DocxDocument

        doc = DocxDocument(str(file_path))
        assets: list[MediaAsset] = []
        seen: set[str] = set()
        context = self._nearest_context(sections, "")

        from docx.opc.constants import RELATIONSHIP_TYPE as RT

        for rel in doc.part.rels.values():
            if rel.reltype != RT.IMAGE:
                continue
            try:
                blob = rel.target_part.blob
                if not blob:
                    continue
                key = hashlib.sha256(blob).hexdigest()
                if key in seen:
                    continue
                seen.add(key)
                partname = rel.target_ref or "image"
                ext = Path(partname).suffix.lower() or ".png"
                mime = MIME_BY_SUFFIX.get(ext, "image/png")
                assets.append(
                    MediaAsset(
                        kind="image",
                        source_path=rel_path,
                        page_no=None,
                        heading_path=f"[图片] DOCX/{Path(partname).name}",
                        mime=mime,
                        data=blob,
                        context_hint=context,
                    )
                )
            except Exception as e:
                logger.debug("DOCX 图片提取失败 {}: {}", rel_path, e)

        return assets

    def _extract_pptx(
        self, file_path: Path, rel_path: str, sections: list[dict]
    ) -> list[MediaAsset]:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        prs = Presentation(str(file_path))
        assets: list[MediaAsset] = []
        seen: set[str] = set()

        for slide_idx, slide in enumerate(prs.slides, start=1):
            context = self._context_for_page(sections, slide_idx)
            img_idx = 0
            for shape in slide.shapes:
                if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                    continue
                img_idx += 1
                try:
                    blob = shape.image.blob
                    if not blob:
                        continue
                    key = hashlib.sha256(blob).hexdigest()
                    if key in seen:
                        continue
                    seen.add(key)
                    ext = shape.image.ext or "png"
                    mime = MIME_BY_SUFFIX.get(f".{ext}", "image/png")
                    assets.append(
                        MediaAsset(
                            kind="image",
                            source_path=rel_path,
                            page_no=slide_idx,
                            heading_path=f"[图片] 幻灯片{slide_idx}/图{img_idx}",
                            mime=mime,
                            data=blob,
                            context_hint=context,
                        )
                    )
                except Exception as e:
                    logger.debug("PPTX 图片提取失败 {} s{}: {}", rel_path, slide_idx, e)
        return assets

    def _collect_local_image_refs(
        self,
        file_path: Path,
        rel_path: str,
        sections: list[dict],
        src_list: list[str],
        *,
        label: str,
    ) -> list[MediaAsset]:
        assets: list[MediaAsset] = []
        seen: set[str] = set()
        context = self._nearest_context(sections, "")
        base_dir = file_path.parent

        for img_idx, src in enumerate(src_list, start=1):
            src = src.strip().split()[0] if src else ""
            if not src or src.startswith("data:"):
                continue
            parsed = urlparse(src)
            if parsed.scheme in ("http", "https"):
                assets.append(
                    MediaAsset(
                        kind="image",
                        source_path=rel_path,
                        page_no=None,
                        heading_path=f"[图片] 外链图{img_idx}",
                        mime="text/plain",
                        data=b"",
                        context_hint=f"外链图片（未下载）: {src}",
                    )
                )
                continue
            local = (base_dir / src).resolve()
            try:
                local.relative_to(self.knowledge_dir)
            except ValueError:
                logger.debug("跳过知识库外图片引用: {} -> {}", rel_path, src)
                continue
            if not local.is_file() or local.suffix.lower() not in IMAGE_SUFFIXES:
                continue
            data = local.read_bytes()
            if not data:
                continue
            key = hashlib.sha256(data).hexdigest()
            if key in seen:
                continue
            seen.add(key)
            assets.append(
                MediaAsset(
                    kind="image",
                    source_path=rel_path,
                    page_no=None,
                    heading_path=f"[图片] {label}/图{img_idx}",
                    mime=guess_mime(local),
                    data=data,
                    context_hint=context or src,
                )
            )
        return assets

    def _extract_markdown_images(
        self, file_path: Path, rel_path: str, sections: list[dict]
    ) -> list[MediaAsset]:
        raw = file_path.read_text(encoding="utf-8", errors="ignore")
        refs = MD_IMAGE_RE.findall(raw)
        refs.extend(HTML_IMG_SRC_RE.findall(raw))
        return self._collect_local_image_refs(
            file_path, rel_path, sections, refs, label="Markdown"
        )

    def _extract_html(
        self, file_path: Path, rel_path: str, sections: list[dict]
    ) -> list[MediaAsset]:
        from bs4 import BeautifulSoup

        raw = file_path.read_text(encoding="utf-8", errors="ignore")
        soup = BeautifulSoup(raw, "html.parser")
        refs = [(tag.get("src") or "").strip() for tag in soup.find_all("img")]
        return self._collect_local_image_refs(
            file_path, rel_path, sections, refs, label="HTML"
        )
